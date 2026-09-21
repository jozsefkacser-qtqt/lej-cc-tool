"""The inspection pick list.

A warehouse document: somebody prints it, walks the racking with it and
ticks parcels off. Two things decide whether it is any good — that it is
ordered the way the walk is, and that a column is never silently blank.
"""

from __future__ import annotations

from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from lej_cc.model import ClearanceStatus, ShipmentRow, Snapshot
from lej_cc.picklist import COLUMNS, build_inspection_list, load_lookup

TZ = ZoneInfo("Europe/Berlin")
NOW = datetime(2026, 9, 18, 14, 30, tzinfo=TZ)


def snap(held: int = 3, cleared: int = 5) -> Snapshot:
    rows = [
        ShipmentRow(
            hawb=f"003404347626045{i:05d}",
            mawb="93602928693",
            status=ClearanceStatus.CLEARED,
            final_status_raw="cleared",
            clearance_time=NOW,
            items=1,
        )
        for i in range(cleared)
    ] + [
        ShipmentRow(
            hawb=f"003404347626099{i:05d}",
            mawb="93602928693",
            status=ClearanceStatus.INSPECTION,
            final_status_raw="not cleared",
            external_status="inspection" if i % 2 else "inspection_doc",
            invoice_number=f"BG-260824{i:05d}",
            atx=f"ATX0343537{i:010d}",
            check_in=NOW,
            items=i + 1,
        )
        for i in range(held)
    ]
    return Snapshot(mawb="93602928693", rows=rows, generated_at=NOW, fetched_at=NOW)


def sheet_of(path):
    return openpyxl.load_workbook(path).active


def rows_of(path) -> list[list]:
    ws = sheet_of(path)
    head = next(n for n, r in enumerate(ws.iter_rows(values_only=True), 1) if r[0] == "Box")
    return [list(r) for r in ws.iter_rows(min_row=head + 1, values_only=True)]


# --- what gets listed ----------------------------------------------------


def test_only_the_parcels_customs_is_holding(tmp_path):
    path = build_inspection_list(snap(held=3, cleared=5), tmp_path)
    assert len(rows_of(path)) == 3


def test_nothing_held_means_no_file(tmp_path):
    """A pick list with nothing to pick is a file somebody opens for nothing."""
    assert build_inspection_list(snap(held=0, cleared=5), tmp_path) is None


def test_the_filename_says_how_many_to_pull(tmp_path):
    path = build_inspection_list(snap(held=7), tmp_path)
    assert path.name == "INSPECTION_936-02928693_7_parcels.xlsx"


def test_no_column_is_silently_blank(tmp_path):
    """An empty column on a printed sheet is worse than no column: somebody
    goes looking for the data that should be in it."""
    path = build_inspection_list(snap(held=3), tmp_path)
    for row in rows_of(path):
        for index, (header, field, _) in enumerate(COLUMNS):
            if field in ("box", "customer", "_tick"):
                continue  # filled from the lookup, or left for a pen
            assert row[index] not in (None, ""), f"{header} is empty"


# --- the walk ------------------------------------------------------------


def test_sorted_by_box_so_the_walk_is_one_pass(tmp_path):
    lookup = {
        "00340434762609900000": {"box": "B-12", "customer": "Acme"},
        "00340434762609900001": {"box": "A-03", "customer": "Bravo"},
        "00340434762609900002": {"box": "B-12", "customer": "Acme"},
    }
    path = build_inspection_list(snap(held=3), tmp_path, lookup)
    boxes = [r[0] for r in rows_of(path)]
    assert boxes == ["A-03", "B-12", "B-12"]


def test_parcels_with_no_box_go_last(tmp_path):
    """They still have to be pulled; they just cannot be walked to."""
    lookup = {"00340434762609900001": {"box": "A-03", "customer": "Acme"}}
    path = build_inspection_list(snap(held=3), tmp_path, lookup)
    boxes = [r[0] for r in rows_of(path)]
    assert boxes[0] == "A-03"
    assert all(b in (None, "") for b in boxes[1:])


def test_the_customer_comes_from_the_lookup(tmp_path):
    lookup = {"00340434762609900000": {"box": "B-1", "customer": "Temu DE"}}
    path = build_inspection_list(snap(held=2), tmp_path, lookup)
    assert "Temu DE" in [r[2] for r in rows_of(path)]


# --- being honest about what PortGround does not give us ----------------


def test_without_a_lookup_the_sheet_says_why_box_is_blank(tmp_path):
    """Inventing a box number sends somebody to the wrong shelf."""
    path = build_inspection_list(snap(held=3), tmp_path)
    notes = [str(c[0].value) for c in sheet_of(path).iter_rows(min_row=1, max_row=3)]
    assert any("PortGround's export carries neither" in n for n in notes)


def test_a_partial_lookup_says_how_many_are_missing(tmp_path):
    lookup = {"00340434762609900000": {"box": "B-1", "customer": "Acme"}}
    path = build_inspection_list(snap(held=3), tmp_path, lookup)
    notes = [str(c[0].value) for c in sheet_of(path).iter_rows(min_row=1, max_row=3)]
    assert any("2 of 3 parcels are not in the lookup" in n for n in notes)


# --- the lookup file -----------------------------------------------------


def test_a_csv_lookup_is_read(tmp_path):
    path = tmp_path / "boxes.csv"
    path.write_text("Tracking number,Box,Customer\n123,A-01,Acme\n", encoding="utf-8")
    assert load_lookup(path) == {"123": {"box": "A-01", "customer": "Acme"}}


def test_an_xlsx_lookup_is_read(tmp_path):
    book = openpyxl.Workbook()
    book.active.append(["HAWB", "Box ID", "Consignee"])
    book.active.append(["123", "A-01", "Acme"])
    path = tmp_path / "boxes.xlsx"
    book.save(path)
    assert load_lookup(path) == {"123": {"box": "A-01", "customer": "Acme"}}


@pytest.mark.parametrize(
    "headers",
    [
        ["Tracking number", "Box", "Customer"],
        ["HAWB", "Box ID", "Consignee"],
        ["barcode", "location", "receiver"],
        ["tracking", "shelf", "name"],
    ],
)
def test_column_names_are_matched_loosely(tmp_path, headers):
    """This file comes from whatever system holds the box numbers. Nobody
    should have to rename headers to use it."""
    path = tmp_path / "boxes.csv"
    path.write_text(",".join(headers) + "\n123,A-01,Acme\n", encoding="utf-8")
    assert load_lookup(path)["123"] == {"box": "A-01", "customer": "Acme"}


def test_a_missing_lookup_is_not_fatal(tmp_path):
    """The pick list is still worth printing without the boxes."""
    assert load_lookup(tmp_path / "nope.csv") == {}
    assert load_lookup(None) == {}


def test_an_unreadable_lookup_is_not_fatal(tmp_path):
    path = tmp_path / "boxes.xlsx"
    path.write_bytes(b"this is not a workbook")
    assert load_lookup(path) == {}


# --- it has to print -----------------------------------------------------


def test_the_sheet_is_set_up_to_print(tmp_path):
    """Landscape, one page wide, header on every sheet — set here rather
    than left to whoever hits Ctrl-P."""
    ws = sheet_of(build_inspection_list(snap(held=60), tmp_path))

    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    # The header row moves with the notes above it, so find it rather than
    # assuming: the point is that these three agree with each other.
    head = next(n for n, r in enumerate(ws.iter_rows(values_only=True), 1) if r[0] == "Box")
    assert ws.print_title_rows == f"${head}:${head}"
    assert ws.freeze_panes == f"A{head + 1}"


def test_there_is_a_column_to_tick(tmp_path):
    ws = sheet_of(build_inspection_list(snap(held=2), tmp_path))
    head = next(n for n, r in enumerate(ws.iter_rows(values_only=True), 1) if r[0] == "Box")
    assert [c.value for c in ws[head]][-1] == "Pulled ✓"


def test_the_repeated_header_is_the_header_and_not_a_blank_line(tmp_path):
    """It only shows up on page two: `append([])` moves the write cursor but
    not max_row, so a predicted header row pointed one line too high."""
    ws = sheet_of(build_inspection_list(snap(held=60), tmp_path))
    head = int(ws.print_title_rows.split(":")[0].lstrip("$"))
    assert [c.value for c in ws[head]][0] == "Box"


# --- the lookup file as it actually arrives -----------------------------


def test_a_hungarian_parcel_list_is_read_without_renaming_anything(tmp_path):
    """QT PARCEL LIST - LEJ is headed in Hungarian; it must work as-is."""
    from lej_cc.picklist import load_lookup

    path = tmp_path / "parcels.csv"
    path.write_text(
        "Csomagszám,Címzett neve,AWB,Karton szám,VÁM Státusz,DEPO,Customer\n"
        "BG-2509165GJ12EV0HU,Ágnes Suszter-Dorogi,235-94755205,HUIS250917587595,"
        "Áruvizsgálatra vár,Express,Radiance Sea Hong Kong Limited\n",
        encoding="utf-8",
    )
    table = load_lookup(path)

    assert table["BG-2509165GJ12EV0HU"]["box"] == "HUIS250917587595"
    assert table["BG-2509165GJ12EV0HU"]["customer"] == "Ágnes Suszter-Dorogi"


def test_the_consignee_wins_over_the_trading_company(tmp_path):
    """A client name repeated down every row does not identify a parcel."""
    from lej_cc.picklist import load_lookup

    path = tmp_path / "both.csv"
    path.write_text(
        "Tracking number,Consignee,Customer,Box\n"
        "TRK1,Ágnes Suszter-Dorogi,Radiance Sea Hong Kong Limited,A-12\n",
        encoding="utf-8",
    )
    assert load_lookup(path)["TRK1"]["customer"] == "Ágnes Suszter-Dorogi"


def test_a_plain_customer_column_is_still_used_when_it_is_all_there_is(tmp_path):
    from lej_cc.picklist import load_lookup

    path = tmp_path / "one.csv"
    path.write_text("Tracking number,Customer,Box\nTRK1,Acme GmbH,A-12\n", encoding="utf-8")
    assert load_lookup(path)["TRK1"]["customer"] == "Acme GmbH"
