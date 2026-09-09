"""Builds the short "what is still open" workbook attached to each update.

PortGround's export is 17 columns wide and lists every shipment under the
master AWB -- 1578 of them on a real one. Nobody chases a shipment from
that. This produces the opposite: only the lines that are still open, only
the columns you need to chase them, oldest first, with a filter row.
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .awb import format_display
from .model import ClearanceStatus, ShipmentRow, Snapshot
from .parser import LOCAL_TZ

log = logging.getLogger(__name__)

#: (header, how to get it from a row, column width). Edit this list to change
#: what the chase sheet contains -- it is the whole definition of the report.
COLUMNS: list[tuple[str, str, int]] = [
    ("HAWB / Tracking number", "hawb", 24),
    ("Status", "final_status_raw", 14),
    ("Days open", "_days_open", 11),
    ("Invoice Number", "invoice_number", 22),
    ("MRN-ID", "mrn_id", 22),
    ("External Status", "external_status", 16),
    ("Check-In", "check_in", 18),
    ("Declaration Sent", "declaration_sent", 18),
    ("Customs Notification", "notification_customs_office", 20),
    ("Items", "items", 8),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OTHER_FILL = PatternFill("solid", fgColor="FFE0E0")


def _days_open(row: ShipmentRow, now: datetime) -> float | None:
    """How long this shipment has been sitting, from the earliest stamp it has.

    None means nothing has happened to it at all yet -- which is its own kind
    of signal, and why those rows sort last rather than first.
    """
    stamps = [s for s in (row.check_in, row.declaration_sent) if s]
    if not stamps:
        return None
    return round((now - min(stamps)).total_seconds() / 86400, 1)


def _sort_key(row: ShipmentRow) -> tuple[int, datetime | None]:
    """Oldest first; shipments with no timestamps at all go last."""
    stamps = [s for s in (row.check_in, row.declaration_sent) if s]
    return (0, min(stamps)) if stamps else (1, None)


def _cell_value(row: ShipmentRow, field: str, now: datetime):  # noqa: ANN201
    if field == "_days_open":
        return _days_open(row, now)
    value = getattr(row, field, None)
    if isinstance(value, datetime):
        return value.astimezone(LOCAL_TZ).strftime("%Y-%m-%d %H:%M")
    return value


def build_open_shipments_workbook(snapshot: Snapshot, dest_dir: Path) -> Path | None:
    """Write the chase sheet for `snapshot`. Returns None if nothing is open."""
    open_rows = snapshot.open_rows
    if not open_rows:
        return None

    now = snapshot.fetched_at or datetime.now(LOCAL_TZ)
    open_rows = sorted(open_rows, key=_sort_key)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Open shipments"

    sheet.append([header for header, _, _ in COLUMNS])
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for row in open_rows:
        sheet.append([_cell_value(row, field, now) for _, field, _ in COLUMNS])
        # An unrecognised status is the one thing a reader must not skim past.
        if row.status is ClearanceStatus.OTHER:
            for index in range(1, len(COLUMNS) + 1):
                sheet.cell(row=sheet.max_row, column=index).fill = OTHER_FILL

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"

    dest_dir.mkdir(parents=True, exist_ok=True)
    path = dest_dir / f"open_{snapshot.mawb}_{now:%Y%m%d_%H%M%S}.xlsx"
    workbook.save(path)
    log.info("wrote chase sheet for %s with %d open rows", snapshot.mawb, len(open_rows))
    return path


def open_shipments_filename(snapshot: Snapshot) -> str:
    """Human-facing name for the attachment in Slack."""
    return (
        f"OPEN_{format_display(snapshot.mawb)}_"
        f"{len(snapshot.open_rows)}_shipments.xlsx"
    )
