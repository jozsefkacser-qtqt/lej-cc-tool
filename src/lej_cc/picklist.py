"""The inspection pick list: which parcels to pull off the shelf.

A warehouse document, not a report. Somebody prints it, walks the floor with
it, and ticks parcels off — so it is landscape, fits the page width, repeats
its header on every sheet, and ends in a column wide enough to write in.

**What PortGround does not give us.** Its export has no box id and no
customer name; the seventeen columns are shipment references and timestamps.
So those two columns are filled from a lookup file you supply, keyed by
tracking number, and left blank with a note on the sheet when there is none.
Blank is the honest state: a box number invented by sorting on something
that is not a box number sends somebody to the wrong shelf.
"""

from __future__ import annotations

import csv
import logging
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .awb import format_display
from .model import ShipmentRow, Snapshot

log = logging.getLogger(__name__)

#: Column header, where the value comes from, width in characters.
COLUMNS: list[tuple[str, str, int]] = [
    ("Box", "box", 14),
    ("Tracking number", "hawb", 24),
    ("Customer", "customer", 26),
    ("Invoice", "invoice_number", 20),
    ("Shipment ref", "atx", 22),
    ("Checked in", "check_in", 17),
    ("Items", "items", 7),
    ("Hold", "external_status", 14),
    ("Pulled ✓", "_tick", 12),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, size=10, color="7A1C1C")
GRID = Border(*[Side(style="thin", color="B8C2CC")] * 4)
#: Alternating bands per box, so a row cannot be read off the wrong line.
BAND = PatternFill("solid", fgColor="EFF3F7")

#: Where a lookup file's columns may be called from, lowercased.
KEY_ALIASES = ("tracking number", "hawb", "hawb / tracking number", "tracking", "barcode")
BOX_ALIASES = ("box", "box id", "box number", "container", "position", "location", "shelf")
CUSTOMER_ALIASES = ("customer", "consignee", "name", "receiver", "recipient")


def _pick(row: dict[str, str], aliases: tuple[str, ...]) -> str:
    for alias in aliases:
        if row.get(alias):
            return str(row[alias]).strip()
    return ""


def load_lookup(path: Path | str | None) -> dict[str, dict[str, str]]:
    """Read tracking number -> {box, customer} from a CSV or XLSX.

    Column names are matched loosely, because this file comes from whatever
    system happens to hold the box numbers and nobody should have to rename
    headers to use it. An unreadable file is logged and ignored: the pick
    list is still worth printing without the boxes.
    """
    if not path:
        return {}
    source = Path(path)
    if not source.exists():
        log.warning("inspection lookup %s does not exist — printing without boxes", source)
        return {}

    try:
        rows: list[dict[str, str]] = []
        if source.suffix.lower() in {".csv", ".tsv", ".txt"}:
            delimiter = "\t" if source.suffix.lower() == ".tsv" else ","
            with source.open(encoding="utf-8-sig", newline="") as handle:
                rows = [
                    {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
                    for row in csv.DictReader(handle, delimiter=delimiter)
                ]
        else:
            book = openpyxl.load_workbook(source, read_only=True, data_only=True)
            sheet = book.active
            lines = sheet.iter_rows(values_only=True)
            header = [str(h or "").strip().lower() for h in next(lines)]
            rows = [
                dict(zip(header, [str(c).strip() if c is not None else "" for c in line],
                         strict=False))
                for line in lines
            ]
            book.close()
    except Exception as exc:  # noqa: BLE001 - never lose the pick list to this
        log.error("could not read the inspection lookup %s: %s", source, exc)
        return {}

    table: dict[str, dict[str, str]] = {}
    for row in rows:
        key = _pick(row, KEY_ALIASES)
        if key:
            table[key] = {
                "box": _pick(row, BOX_ALIASES),
                "customer": _pick(row, CUSTOMER_ALIASES),
            }
    log.info("inspection lookup %s: %d tracking numbers", source, len(table))
    return table


@dataclass(frozen=True)
class Pick:
    """One parcel to pull, and where it is if we know."""

    row: ShipmentRow
    box: str
    customer: str

    @property
    def hawb(self) -> str:
        return self.row.hawb

    def sort_key(self) -> tuple[int, str, str]:
        """By box, so the walk is one pass of the racking; no box sorts last."""
        return (0, self.box.casefold(), self.hawb) if self.box else (1, "", self.hawb)


def _value(pick: Pick, field: str) -> object:
    if field == "box":
        return pick.box
    if field == "customer":
        return pick.customer
    if field == "_tick":
        return ""
    value = getattr(pick.row, field, None)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return value


def build_inspection_list(
    snapshot: Snapshot, dest_dir: Path, lookup: dict[str, dict[str, str]] | None = None
) -> Path | None:
    """Write the pick list for `snapshot`. None when nothing is held."""
    held = snapshot.inspection_rows
    if not held:
        return None

    lookup = lookup or {}
    entries = [
        Pick(
            row=row,
            box=lookup.get(row.hawb, {}).get("box", ""),
            customer=lookup.get(row.hawb, {}).get("customer", ""),
        )
        for row in held
    ]
    entries.sort(key=Pick.sort_key)
    matched = sum(1 for e in entries if e.box)

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Inspection pick list"
    mawb = format_display(snapshot.mawb)
    now = snapshot.fetched_at or datetime.now()

    sheet.append([f"Customs inspection — {mawb}"])
    sheet["A1"].font = TITLE_FONT
    sheet.append([
        f"{len(held)} parcel(s) to pull · data from PortGround "
        f"{now:%Y-%m-%d %H:%M} · printed from the LEJ customs tracker"
    ])
    if not matched:
        sheet.append([
            "Box and Customer are blank: PortGround's export carries neither. "
            "Point INSPECTION_LOOKUP_FILE at a file keyed by tracking number to fill them."
        ])
        sheet.cell(row=3, column=1).font = NOTE_FONT
    elif matched < len(entries):
        sheet.append([
            f"{len(entries) - matched} of {len(entries)} parcels are not in the lookup "
            "file — they are listed last, without a box."
        ])
        sheet.cell(row=3, column=1).font = NOTE_FONT
    sheet.append([])

    # Read the row back after writing it rather than predicting it: an
    # `append([])` moves the write cursor but not max_row, so a predicted
    # number pointed at the blank line above -- and the header repeated on
    # page two onwards would have been that blank line.
    sheet.append([header for header, _, _ in COLUMNS])
    head_row = sheet.max_row
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=head_row, column=index)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")
    sheet.row_dimensions[head_row].height = 22

    previous_box, band = None, False
    for entry in entries:
        if entry.box != previous_box:
            band, previous_box = not band, entry.box
        sheet.append([_value(entry, field) for _, field, _ in COLUMNS])
        for index in range(1, len(COLUMNS) + 1):
            cell = sheet.cell(row=sheet.max_row, column=index)
            cell.border = GRID
            if band:
                cell.fill = BAND
        sheet.row_dimensions[sheet.max_row].height = 20

    # Printing is the point, so set it up rather than leaving it to whoever
    # hits Ctrl-P: landscape, one page wide, header repeated on every sheet.
    sheet.freeze_panes = sheet.cell(row=head_row + 1, column=1)
    sheet.print_title_rows = f"{head_row}:{head_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.gridLines = False
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.oddFooter.left.text = f"{mawb} — inspection pick list"

    dest_dir.mkdir(parents=True, exist_ok=True)
    path = dest_dir / f"INSPECTION_{mawb}_{len(held)}_parcels.xlsx"
    book.save(path)
    log.info(
        "wrote the inspection pick list for %s: %d parcel(s), %d with a box",
        snapshot.mawb, len(held), matched,
    )
    return path
